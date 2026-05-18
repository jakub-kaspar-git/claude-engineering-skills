"""
Sdílený toolkit pro generování .docx dokumentů se státnicovými otázkami.

Poskytuje OMML rovnice, formátované nadpisy, tabulky, info/warning boxy,
obrázky s popisky, stránkování a PDF export.

Použití v per-question skriptu:
    from docx_engine import (
        create_document, save_and_export, add_heading, add_para,
        add_equation, add_bullet, add_image, add_info_box,
        add_warning_box, add_page_break, add_styled_table,
        COLORS, IMG_DIR,
    )
"""

import os
from datetime import date
from lxml import etree
import latex2mathml.converter
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import nsdecls
from docx.oxml import parse_xml

# ── Paths ────────────────────────────────────────────────────────────────────

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_DIR = os.path.join(BASE_DIR, "output")
IMG_DIR = os.path.join(OUTPUT_DIR, "img")
os.makedirs(IMG_DIR, exist_ok=True)

# ── OMML helpers ─────────────────────────────────────────────────────────────

OMML_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"

XSLT_PATH = None
for _candidate in [
    os.path.expandvars(r"%ProgramFiles%\Microsoft Office\root\Office16\MML2OMML.XSL"),
    os.path.expandvars(r"%ProgramFiles(x86)%\Microsoft Office\root\Office16\MML2OMML.XSL"),
    os.path.expandvars(r"%ProgramFiles%\Microsoft Office\Office16\MML2OMML.XSL"),
    os.path.expandvars(r"%ProgramFiles(x86)%\Microsoft Office\Office16\MML2OMML.XSL"),
    os.path.expandvars(r"%ProgramFiles%\Microsoft Office\root\Office15\MML2OMML.XSL"),
]:
    if os.path.isfile(_candidate):
        XSLT_PATH = _candidate
        break

_transform = None


def get_transform():
    """Get (or lazily create) the MathML→OMML XSLT transform."""
    global _transform
    if _transform is None:
        xslt_tree = etree.parse(XSLT_PATH)
        _transform = etree.XSLT(xslt_tree)
    return _transform


def latex_to_omml(latex_str: str) -> etree._Element:
    """Convert a LaTeX math string to an OMML element for python-docx."""
    mathml_str = latex2mathml.converter.convert(latex_str)
    mathml_tree = etree.fromstring(mathml_str.encode("utf-8"))
    omml_tree = get_transform()(mathml_tree)
    omml_root = omml_tree.getroot()
    omath = omml_root.find(f"{{{OMML_NS}}}oMath")
    return omath if omath is not None else omml_root


# ── Pagination helpers ───────────────────────────────────────────────────────

def _keep_with_next(para):
    """Set 'keep with next' so this paragraph stays on the same page as the following one."""
    pPr = para._element.get_or_add_pPr()
    pPr.append(parse_xml(f'<w:keepNext {nsdecls("w")}/>'))
    return para


def _keep_together(para):
    """Prevent this paragraph from being split across pages."""
    pPr = para._element.get_or_add_pPr()
    pPr.append(parse_xml(f'<w:keepLines {nsdecls("w")}/>'))
    return para


# ── Document helpers ─────────────────────────────────────────────────────────

def add_page_break(doc):
    """Insert an explicit page break."""
    para = doc.add_paragraph()
    run = para.add_run()
    run.add_break(WD_BREAK.PAGE)
    return para


def add_equation(doc, latex_str, label=None):
    """Insert a centered OMML equation with optional label."""
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(4)
    para.paragraph_format.space_after = Pt(4)
    _keep_together(para)
    _keep_with_next(para)
    para._element.append(latex_to_omml(latex_str))
    if label:
        run = para.add_run(f"   ({label})")
        run.font.size = Pt(10)
        run.font.italic = True
    return para


def add_heading(doc, text, level=1):
    """Add a styled heading that stays with the next paragraph."""
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)
    _keep_with_next(h)
    return h


def add_para(doc, text, bold=False, italic=False, size=11, keep_with_next=False):
    """Add a paragraph with formatting options."""
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    _keep_together(para)
    if keep_with_next:
        _keep_with_next(para)
    return para


def add_bullet(doc, text, level=0, is_last=False):
    """Add a bullet point. All bullets except the last in a group get keep_with_next."""
    para = doc.add_paragraph(text, style="List Bullet")
    if level > 0:
        para.paragraph_format.left_indent = Cm(1.27 * (level + 1))
    _keep_together(para)
    if not is_last:
        _keep_with_next(para)
    return para


def add_image(doc, path, width_cm=14, caption=None):
    """Add an image centered with optional caption. Image + caption stay together."""
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _keep_together(para)
    if caption:
        _keep_with_next(para)
    run = para.add_run()
    run.add_picture(path, width=Cm(width_cm))
    if caption:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _keep_together(cap)
        r = cap.add_run(caption)
        r.font.size = Pt(9)
        r.font.italic = True
        r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    return para


def add_info_box(doc, title, text):
    """Add a highlighted info box (light blue) using a 1-cell table with shading."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = table._element.tblPr
    tblPr.append(parse_xml(f'<w:tblpPr {nsdecls("w")}/>'))
    cell = table.rows[0].cells[0]
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="E8F0FE"/>')
    cell._element.get_or_add_tcPr().append(shading)
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
    # Title
    p = cell.paragraphs[0]
    _keep_together(p)
    _keep_with_next(p)
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(0x1A, 0x47, 0x8A)
    # Body
    p2 = cell.add_paragraph()
    _keep_together(p2)
    r2 = p2.add_run(text)
    r2.font.size = Pt(10)
    doc.add_paragraph()  # spacer


def add_warning_box(doc, text):
    """Add a yellow warning box."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="FFF3CD"/>')
    cell._element.get_or_add_tcPr().append(shading)
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
    p = cell.paragraphs[0]
    _keep_together(p)
    r = p.add_run("Pozor: ")
    r.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(0x85, 0x6A, 0x04)
    r2 = p.add_run(text)
    r2.font.size = Pt(10)
    doc.add_paragraph()


def add_exam_questions(doc, questions: list):
    """Add exam Q&A section. Each item is a (question, answer) tuple."""
    add_page_break(doc)
    add_heading(doc, "Část 4: Typické zkouškové otázky", level=2)
    add_para(doc,
        "Otázky, které mohou padnout u ústní státní závěrečné zkoušky, "
        "s doporučenými odpověďmi.", italic=True, size=10)
    doc.add_paragraph()

    for i, (q, a) in enumerate(questions, start=1):
        # Question as green box
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.rows[0].cells[0]
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="DCFCE7"/>')
        cell._element.get_or_add_tcPr().append(shading)
        for row in table.rows:
            trPr = row._tr.get_or_add_trPr()
            trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        p = cell.paragraphs[0]
        _keep_together(p)
        r = p.add_run(f"Otázka {i}: ")
        r.bold = True
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor(0x16, 0x6A, 0x34)
        r2 = p.add_run(q)
        r2.bold = True
        r2.font.size = Pt(11)
        r2.font.color.rgb = RGBColor(0x16, 0x6A, 0x34)

        # Answer
        ans = doc.add_paragraph()
        _keep_together(ans)
        r_label = ans.add_run("Odpověď: ")
        r_label.bold = True
        r_label.font.size = Pt(10)
        r_body = ans.add_run(a)
        r_body.font.size = Pt(10)
        doc.add_paragraph()  # spacer


def add_styled_table(doc, headers: list, data: list, style="Light Grid Accent 1"):
    """Add a table with bold header row and consistent formatting."""
    table = doc.add_table(rows=1 + len(data), cols=len(headers))
    table.style = style
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # Header row
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
    # Data rows
    for r_idx, row_data in enumerate(data, start=1):
        for c_idx, val in enumerate(row_data):
            table.rows[r_idx].cells[c_idx].text = val
    doc.add_paragraph()  # spacer
    return table


# ── Document lifecycle ───────────────────────────────────────────────────────

def create_document(title: str, okruh: str) -> Document:
    """Create a new Document with standard styling and header."""
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    add_para(doc, "Státnicová otázka", bold=True, size=16)
    add_para(doc, f"Okruh: {okruh}  |  Datum: {date.today().strftime('%d. %m. %Y')}",
             size=10, italic=True)
    doc.add_paragraph()
    add_heading(doc, title, level=1)
    return doc


def save_and_export(doc, filename: str) -> tuple:
    """Save .docx and export PDF. Returns (docx_path, pdf_path_or_None)."""
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    docx_path = os.path.join(OUTPUT_DIR, f"{filename}.docx")
    doc.save(docx_path)
    print(f"Dokument uložen: {docx_path}")

    pdf_dir = os.path.join(OUTPUT_DIR, "pdf")
    os.makedirs(pdf_dir, exist_ok=True)
    pdf_path = os.path.join(pdf_dir, f"{filename}.pdf")
    try:
        docx_to_pdf(docx_path, pdf_path)
        print(f"PDF uložen: {pdf_path}")
        return docx_path, pdf_path
    except Exception as e:
        print(f"Varování: PDF export selhal ({e}). Soubor .docx byl uložen.")
        return docx_path, None


# ── PDF export ───────────────────────────────────────────────────────────────

def docx_to_pdf(docx_path, pdf_path):
    """Convert .docx to .pdf using Microsoft Word COM automation."""
    import comtypes.client
    word = comtypes.client.CreateObject("Word.Application")
    word.Visible = False
    try:
        doc = word.Documents.Open(os.path.abspath(docx_path))
        doc.SaveAs(os.path.abspath(pdf_path), FileFormat=17)  # 17 = wdFormatPDF
        doc.Close()
    finally:
        word.Quit()


# ── Matplotlib defaults ──────────────────────────────────────────────────────

plt.rcParams.update({
    "font.family": "sans-serif",
    "font.size": 10,
    "axes.grid": True,
    "grid.alpha": 0.3,
    "figure.dpi": 200,
})

COLORS = ["#2563EB", "#DC2626", "#16A34A", "#D97706", "#7C3AED", "#0891B2"]
